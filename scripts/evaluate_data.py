"""Replay transformed public retail invoice amounts through the PostgreSQL ledger."""

import argparse
import hashlib
import io
import json
import platform
import time
import zipfile
from datetime import datetime, timezone
from decimal import ROUND_HALF_UP, Decimal
from uuid import uuid4

import httpx
from openpyxl import load_workbook

from veriforge.config import ROOT
from veriforge.db import connect, seed_accounts
from veriforge.finance import transfer

URL = "https://archive.ics.uci.edu/static/public/352/online+retail.zip"


def main(limit):
    folder = ROOT / ".local/data"
    folder.mkdir(parents=True, exist_ok=True)
    file = folder / "online-retail.zip"
    if not file.exists():
        with httpx.stream("GET", URL, follow_redirects=True, timeout=120) as response:
            response.raise_for_status()
            with file.open("wb") as f:
                for block in response.iter_bytes():
                    f.write(block)
    checksum = hashlib.sha256(file.read_bytes()).hexdigest()
    with zipfile.ZipFile(file) as archive:
        member = next(n for n in archive.namelist() if n.endswith(".xlsx"))
        workbook = load_workbook(io.BytesIO(archive.read(member)), read_only=True, data_only=True)
    rows = workbook.active.iter_rows(values_only=True)
    header = next(rows)
    idx = {name: i for i, name in enumerate(header)}
    tenant = "dataset-" + uuid4().hex[:12]
    seed_accounts(tenant)
    durations = []
    skipped = 0
    failures = []
    for row_number, row in enumerate(rows, 2):
        invoice = str(row[idx["InvoiceNo"]])
        quantity = row[idx["Quantity"]]
        price = row[idx["UnitPrice"]]
        if invoice.startswith("C") or not price or quantity <= 0 or price <= 0:
            skipped += 1
            continue
        amount = (Decimal(str(price)) * quantity).quantize(Decimal(".01"), rounding=ROUND_HALF_UP)
        if amount > 1000:
            skipped += 1
            continue
        start = time.monotonic()
        key = f"uci-row-{row_number}"
        try:
            first = transfer(tenant, "alice", "bob", str(amount), "THB", key)
            second = transfer(tenant, "alice", "bob", str(amount), "THB", key)
            assert first["id"] == second["id"] and second["replayed"]
        except Exception as e:
            failures.append({"row_number": row_number, "error": str(e)[:200]})
        durations.append((time.monotonic() - start) * 1000)
        if len(durations) >= limit:
            break
    workbook.close()
    with connect() as conn:
        actual = conn.execute("SELECT count(*) n FROM payments WHERE tenant=%s", (tenant,)).fetchone()["n"]
        imbalance = conn.execute(
            "SELECT count(*) n FROM (SELECT p.id FROM payments p JOIN postings l ON l.payment_id=p.id WHERE p.tenant=%s GROUP BY p.id HAVING sum(l.delta)<>0) x",
            (tenant,),
        ).fetchone()["n"]
    report = {
        "dataset": "UCI Online Retail",
        "source": URL,
        "license": "CC BY 4.0 (UCI dataset page)",
        "sha256": checksum,
        "retrieved_at": datetime.now(timezone.utc).isoformat(),
        "transform_version": "retail-replay-v1",
        "source_provenance": "PUBLIC_BENCHMARK",
        "execution_provenance": "SYNTHETIC_REPLAY",
        "transform": "Positive non-cancelled line-item quantity × GBP unit price, <=1000, rounded half-up; numerical amounts mapped to THB sandbox units without FX conversion. Not bank transfers or historical account reconstruction.",
        "sampling": "first eligible rows in source order; no random sampling",
        "excluded_rows": skipped,
        "samples": len(durations),
        "payments_posted": actual,
        "failures": failures,
        "ledger_imbalances": imbalance,
        "p95_two_request_ms": sorted(durations)[max(0, int(len(durations) * 0.95) - 1)],
        "runtime": platform.python_version(),
        "platform": platform.platform(),
        "tenant": tenant,
        "limitations": [
            "One local machine",
            "Sequential workload",
            "No real customer identifiers retained",
            "Not a fraud benchmark",
        ],
    }
    out = ROOT / "artifacts"
    out.mkdir(exist_ok=True)
    (out / "dataset-evaluation.json").write_text(json.dumps(report, indent=2))
    print(
        json.dumps(
            {
                "samples": len(durations),
                "payments": actual,
                "failures": len(failures),
                "imbalances": imbalance,
            }
        )
    )
    if failures or imbalance or actual != len(durations):
        raise SystemExit(1)


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--limit", type=int, default=100)
    args = parser.parse_args()
    if not 1 <= args.limit <= 1000:
        parser.error("limit must be 1..1000")
    main(args.limit)
